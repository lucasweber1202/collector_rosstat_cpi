"""Network and file boundary tests; fixtures are synthetic, not Rosstat layouts."""

from __future__ import annotations

import io
import zipfile

import httpx
import pytest
from openpyxl import Workbook

from scripts import source_research as source


@pytest.mark.parametrize(
    "url",
    [
        "http://rosstat.gov.ru/data",
        "https://rosstat.gov.ru.evil.test/data",
        "https://evil.test/rosstat.gov.ru",
        "https://user:password@rosstat.gov.ru/data",
        "https://rosstat.gov.ru:8443/data",
        "https://127.0.0.1/data",
        "https://rosstat.gov.ru/data?token=secret",
    ],
)
def test_reject_untrusted_urls(url: str) -> None:
    with pytest.raises(ValueError):
        source.validate_url(url)


def test_retry_then_success(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = []
    sleeps = []

    def respond(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(
            503 if len(calls) == 1 else 200,
            headers={"content-type": "text/html"},
            text="<html>ok</html>",
        )

    monkeypatch.setattr(source.time, "sleep", sleeps.append)
    with httpx.Client(transport=httpx.MockTransport(respond)) as client:
        result = source.fetch_document(client, source.SOURCE_URL)
    assert len(calls) == 2 and len(sleeps) == 1
    assert result["body"] == b"<html>ok</html>"


def test_redirect_never_contacts_external_host() -> None:
    calls = []

    def respond(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(302, headers={"location": "https://evil.test/data"})

    with httpx.Client(transport=httpx.MockTransport(respond)) as client, pytest.raises(ValueError):
        source.fetch_document(client, source.SOURCE_URL)
    assert len(calls) == 1


def test_bounded_redirect_loop() -> None:
    with (
        httpx.Client(
            transport=httpx.MockTransport(
                lambda req: httpx.Response(302, headers={"location": source.SOURCE_URL})
            )
        ) as client,
        pytest.raises(ValueError, match="redirect"),
    ):
        source.fetch_document(client, source.SOURCE_URL)


@pytest.mark.parametrize("status", [401, 403, 404])
def test_nonretryable_http_error(status: int) -> None:
    calls = []

    def respond(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(status, text="do not print response cookies or body")

    with (
        httpx.Client(transport=httpx.MockTransport(respond)) as client,
        pytest.raises(RuntimeError, match=str(status)),
    ):
        source.fetch_document(client, source.SOURCE_URL)
    assert len(calls) == 1


def test_size_limit_before_parsing(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(source, "MAX_DOWNLOAD_BYTES", 8)
    with (
        httpx.Client(
            transport=httpx.MockTransport(
                lambda req: httpx.Response(
                    200, headers={"content-type": "text/html"}, content=b"x" * 9
                )
            )
        ) as client,
        pytest.raises(ValueError, match="size"),
    ):
        source.fetch_document(client, source.SOURCE_URL)


def test_discover_preserves_multiple_candidates_and_excludes_regions() -> None:
    html = """<a href="/storage/mediabank/a.xlsx">Индексы потребительских цен в группировке КИПЦ</a>
    <a href='/storage/mediabank/b.xlsx'><span>Индексы потребительских цен в группировке КИПЦ</span></a>
    <a href='/storage/mediabank/c.xlsx'>Индексы потребительских цен по областям</a>
    <a href='https://evil.test/d.xlsx'>КИПЦ</a>"""
    links = source.discover_links(html, source.SOURCE_URL)
    assert {row["url"].rsplit("/", 1)[-1] for row in links} == {"a.xlsx", "b.xlsx"}


def test_workbook_inspection_preserves_cells_without_interpreting_them() -> None:
    book = Workbook()
    book.active.append(["КИПЦ", "Январь 2026", "Вес"])
    book.active.append(["01.1", 100.25, 12.5])
    buf = io.BytesIO()
    book.save(buf)
    result = source.inspect_document(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert result["kind"] == "xlsx"
    assert result["sheets"][0]["sample_rows"][1] == ["01.1", 100.25, 12.5]


def test_html_disguised_as_xlsx_fails() -> None:
    with pytest.raises(ValueError, match="signature"):
        source.inspect_document(
            b"<html>Error</html>",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


def test_zip_bomb_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "x" * 10000)
    monkeypatch.setattr(source, "MAX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(ValueError, match="uncompressed"):
        source.inspect_document(buf.getvalue(), "application/zip")


def test_transport_errors_do_not_expose_proxy_secrets(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(source, "MAX_RETRIES", 0)

    def fail(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("proxy password=TOPSECRET", request=request)

    with (
        httpx.Client(transport=httpx.MockTransport(fail)) as client,
        pytest.raises(RuntimeError) as exc,
    ):
        source.fetch_document(client, source.SOURCE_URL)
    assert "TOPSECRET" not in str(exc.value)


def test_partial_acquisition_keeps_success_and_reports_failure(tmp_path, monkeypatch) -> None:
    page = source.SOURCE_URL
    failed = "https://rosstat.gov.ru/missing.html"
    book_url = "https://rosstat.gov.ru/storage/mediabank/basket.xlsx"
    book = Workbook()
    book.active.append(["КИПЦ", "Вес"])
    blob = io.BytesIO()
    book.save(blob)
    html = '<a href="/storage/mediabank/basket.xlsx">Структура потребительских расходов</a>'
    calls = []

    def respond(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        calls.append(url)
        if url == failed:
            return httpx.Response(404)
        if url == page:
            return httpx.Response(
                200,
                headers={"content-type": "text/html; charset=windows-1251"},
                content=html.encode("cp1251"),
            )
        assert url == book_url
        return httpx.Response(
            200, headers={"content-type": source.XLSX_MIME}, content=blob.getvalue()
        )

    monkeypatch.setattr(source, "SOURCE_PAGES", (page, failed))
    monkeypatch.setattr(
        source, "_build_client", lambda: httpx.Client(transport=httpx.MockTransport(respond))
    )
    report = source.research_sources(tmp_path)
    assert report["downloaded_workbooks"] == 1
    assert len(report["errors"]) == 1
    assert report["statistical_validation"] == "NOT_PERFORMED"
    assert calls == [page, failed, book_url]
    assert (tmp_path / report["documents"][1]["file"]).read_bytes() == blob.getvalue()


def test_client_keeps_tls_verification_and_environment_proxy(monkeypatch) -> None:
    settings = {}

    def client(**kwargs):
        settings.update(kwargs)

    monkeypatch.setattr(source.httpx, "Client", client)
    source._build_client()
    assert settings["verify"] is True
    assert settings["trust_env"] is True
    assert settings["follow_redirects"] is False


def test_xml_security_cannot_be_disabled(monkeypatch) -> None:
    book = Workbook()
    buf = io.BytesIO()
    book.save(buf)
    monkeypatch.setattr(source, "DEFUSEDXML", False)
    with pytest.raises(RuntimeError, match="defusedxml"):
        source.inspect_document(buf.getvalue(), source.XLSX_MIME)
