"""Download and parse the official Rosstat consumer price publications.

Rosstat publishes the CPI as a set of workbooks under one mediabank path plus an
SDMX 2.1 message prepared for the IMF SDDS. Four surfaces are used:

``ipc_spr_{MM}-{YYYY}.xlsx``
    "Индексы потребительских цен и структура потребительских расходов по
    Российской Федерации" - one sheet per reference month carrying, for every
    published node of the Rosstat CPI grouping, the consumer-expenditure weight
    and the index against the previous month, against December of the previous
    year and against the same month of the previous year.
``ipc_mes_{MM}-{YYYY}.xlsx``
    The same month-on-month index for the four headline aggregates back to 1991,
    which is the only surface carrying the pre-2025 monthly history.
``CPR_tov_RF_*.xlsx`` / ``CPR_Gr-Tov-RF_*.xlsx``
    The annual consumer-expenditure structure by local code, which extends the
    official basket back beyond the combined publication.
``SDDS_CPI*.xml``
    SDMX 2.1 StructureSpecificData against the IMF ECOFIN data structure. It is
    the only Rosstat surface publishing headline CPI as an index *level* rather
    than as a percentage change, and it carries two index reference periods.

``Vesa-tov-KIPC_*.xlsx`` carries the official KIPC (COICOP) classification with
its hierarchical codes; it is read for hierarchy and audit output only, because
Rosstat publishes KIPC indices annually and this collector is monthly.
"""

from __future__ import annotations

import io
import logging
import math
import re
import ssl
import time
from datetime import date
from typing import Any, Final, NamedTuple
from urllib.parse import urljoin, urlsplit

import httpx
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from scripts.config import (
    BACKOFF_FACTOR,
    CA_BUNDLE_PATH,
    DOWNLOAD_DELAY,
    MAX_DOWNLOAD_BYTES,
    MAX_RETRIES,
    MAX_RETRY_DELAY,
    RATE_LIMIT_BACKOFF,
    REQUEST_TIMEOUT,
    USER_AGENT,
)

logger = logging.getLogger(__name__)

# --------------------------------------------------------------------------
# Verified source constants
# --------------------------------------------------------------------------

PRIMARY_HOST: Final = "https://rosstat.gov.ru"
# eng.rosstat.gov.ru serves the identical /storage/mediabank/ objects and is
# used when the primary host fails; both were byte-compared during research.
FALLBACK_HOST: Final = "https://eng.rosstat.gov.ru"
HOSTS: Final = (PRIMARY_HOST, FALLBACK_HOST)
# Server-side requests never leave this allowlist, whatever a discovered page
# happens to link to.
ALLOWED_HOSTS: Final = frozenset({"rosstat.gov.ru", "eng.rosstat.gov.ru"})

PRICE_SECTION_PATH: Final = "/statistics/price"
SDMX_SECTION_PATH: Final = "/folder/210404"
MEDIABANK_PREFIX: Final = "/storage/mediabank/"

SOURCE_NAME: Final = "Federal State Statistics Service (Rosstat)"
RELEASE_NAME: Final = "Consumer Price Index and Structure of Consumer Expenditures"
# The official title of each surface, used verbatim in stored descriptions.
RELEASE_NAMES: Final = {
    "monthly_detail": (
        "Rosstat, Consumer price indices and the structure of consumer expenditures "
        "for the Russian Federation"
    ),
    "monthly_headline": (
        "Rosstat, Consumer price indices for goods and services in the Russian Federation, "
        "monthly since 1991"
    ),
    "sdmx": "Rosstat, SDMX SDDS consumer price index (IMF ECOFIN_DSD 1.0)",
}
PRICE_SECTION_URL: Final = f"{PRIMARY_HOST}{PRICE_SECTION_PATH}"
SDMX_SECTION_URL: Final = f"{FALLBACK_HOST}{SDMX_SECTION_PATH}"
METHODOLOGY_URL: Final = f"{PRIMARY_HOST}/storage/mediabank/Opredeleniya_IPC.pdf"

# Discovery patterns. The monthly files carry the reference month in their name,
# so the current names are read from the official section page instead of being
# guessed or pinned.
_FILE_PATTERNS: Final = {
    "monthly_detail": re.compile(r"ipc_spr_\d{2}-\d{4}\.xlsx"),
    "monthly_headline": re.compile(r"ipc_mes_\d{2}-\d{4}\.xlsx"),
    "item_weights": re.compile(r"CPR_tov_RF_\d{4}-\d{4}\.xlsx"),
    "group_weights": re.compile(r"CPR_Gr-Tov-RF_\d{4}-\d{4}\.xlsx"),
    "kipc_weights": re.compile(r"Vesa-tov-KIPC_\d{4}-\d{4}\.xlsx"),
}
_SDMX_PATTERNS: Final = {
    "sdmx_index_2000": re.compile(r"SDDS_CPI[ _]2000_\d{4}\.xml"),
    "sdmx_index_2010": re.compile(r"SDDS_CPI_PPI_\d{4}\.xml"),
}

RUSSIAN_MONTHS: Final = {
    "январь": 1,
    "января": 1,
    "февраль": 2,
    "февраля": 2,
    "март": 3,
    "марта": 3,
    "апрель": 4,
    "апреля": 4,
    "май": 5,
    "мая": 5,
    "июнь": 6,
    "июня": 6,
    "июль": 7,
    "июля": 7,
    "август": 8,
    "августа": 8,
    "сентябрь": 9,
    "сентября": 9,
    "октябрь": 10,
    "октября": 10,
    "ноябрь": 11,
    "ноября": 11,
    "декабрь": 12,
    "декабря": 12,
}

# --------------------------------------------------------------------------
# Structured identifiers
# --------------------------------------------------------------------------

MEASURE_PREFIX: Final = "CPI"
GEO: Final = "RU"
CLASSIFICATIONS: Final = frozenset({"RSTG", "SDDS"})
LEVELS: Final = frozenset({"HEADLINE", "AGGREGATE", "GROUP", "ITEM"})
MEASURES: Final = frozenset({"MOM", "YTD", "YOY", "IX2000", "IX2010"})

COUNTRY_CURRENCY: Final = "RUB"
FREQUENCIES: Final = frozenset({"monthly"})
UNITS: Final = frozenset({"index"})
ECO_GROUPS: Final = frozenset({"consumer_prices"})

# Rosstat uses a Cyrillic suffix on the pharmaceutical ATC codes it publishes in
# the monthly workbook (for example ``8060.АГ``). Identifiers stay ASCII, so the
# two letters that actually occur are transliterated deterministically.
_NODE_TRANSLITERATION: Final = {"А": "A", "Г": "G"}
_NODE_ALLOWED = re.compile(r"^[A-Z0-9-]+$")

# The Rosstat CPI grouping publishes the headline, the two goods/services legs
# and a set of analytical cuts that deliberately overlap each other (core CPI,
# baskets excluding alcohol or excluding vegetables). Only the first group forms
# the published aggregation tree; the rest are alternative aggregates over the
# same basket and must never be summed together.
HEADLINE_CODE: Final = "1"
TREE_PARENTS: Final = {"1": ("2", "9000"), "2": ("6", "7")}
ANALYTICAL_CODES: Final = frozenset({"3", "4", "5", "70", "71", "72", "80"})
AGGREGATE_CODES: Final = frozenset({"2", "6", "7", "9000"}) | ANALYTICAL_CODES
# ipc_mes sheet title fragment -> local code of the aggregate it publishes.
HEADLINE_SHEET_CONCEPTS: Final = (
    ("на товары и услуги", "1"),
    ("на продовольственные товары", "6"),
    ("на непродовольственные товары", "7"),
    ("на услуги", "9000"),
)

# SDMX indicator/base pairs actually present in the published messages.
SDMX_MEASURES: Final = {"sdmx_index_2000": "IX2000", "sdmx_index_2010": "IX2010"}
SDMX_BASE_PERIODS: Final = {"IX2000": "2000=100", "IX2010": "2010=100"}


class SeriesKey(NamedTuple):
    """Decoded parts of a ``series_id``, ordered coarse to fine."""

    measure_prefix: str
    geo: str
    classification: str
    level: str
    node: str
    measure: str


def normalise_node(code: str) -> str:
    """Return the ASCII, uppercase identifier part for a native Rosstat code."""
    text = str(code).strip().upper()
    for cyrillic, latin in _NODE_TRANSLITERATION.items():
        text = text.replace(cyrillic, latin)
    text = re.sub(r"[.\s_/]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-")
    if not _NODE_ALLOWED.match(text):
        raise ValueError(f"Rosstat code is not representable in a series_id: {code!r}")
    return text


def build_series_id(classification: str, level: str, node: str, measure: str) -> str:
    """Compose the structured identifier for one published series."""
    if classification not in CLASSIFICATIONS:
        raise ValueError(f"Unknown classification: {classification}")
    if level not in LEVELS:
        raise ValueError(f"Unknown level: {level}")
    if measure not in MEASURES:
        raise ValueError(f"Unknown measure: {measure}")
    return "_".join((MEASURE_PREFIX, GEO, classification, level, normalise_node(node), measure))


def parse_series_id(series_id: str) -> SeriesKey:
    """Decode ``CPI_RU_{classification}_{level}_{node}_{measure}``."""
    parts = series_id.split("_")
    if len(parts) != 6:
        raise ValueError(f"Invalid Rosstat CPI series_id: {series_id}")
    key = SeriesKey(*parts)
    if key.measure_prefix != MEASURE_PREFIX or key.geo != GEO:
        raise ValueError(f"Invalid Rosstat CPI series_id: {series_id}")
    if key.classification not in CLASSIFICATIONS or key.level not in LEVELS:
        raise ValueError(f"Invalid Rosstat CPI series_id: {series_id}")
    if key.measure not in MEASURES or not _NODE_ALLOWED.match(key.node):
        raise ValueError(f"Invalid Rosstat CPI series_id: {series_id}")
    return key


# --------------------------------------------------------------------------
# Upstream descriptive metadata gathered during extraction
# --------------------------------------------------------------------------

# Populated by collect_raw_data() and consumed by metadata.py. Cleared on entry
# to every collection so a rerun never reuses a previous run's descriptions.
UPSTREAM_METADATA: dict[str, dict[str, Any]] = {}
_SOURCE_FILES: dict[str, str] = {}
_LAST_PUBLISH_DATE: dict[str, date] = {}
_HIERARCHY: dict[str, dict[str, Any]] = {}
_KIPC_NODES: dict[int, dict[str, dict[str, Any]]] = {}


def get_series_catalog() -> dict[str, dict[str, Any]]:
    """Return the upstream descriptive row per series seen in this collection."""
    return UPSTREAM_METADATA


def get_source_files() -> dict[str, str]:
    """Return the resolved absolute URL of every source file this run read."""
    return dict(_SOURCE_FILES)


def get_last_publish_date() -> date | None:
    """Return the newest source-stamped publication date seen in this run."""
    return max(_LAST_PUBLISH_DATE.values()) if _LAST_PUBLISH_DATE else None


def get_hierarchy() -> dict[str, dict[str, Any]]:
    """Return one node row per published CPI node, keyed by local code."""
    return _HIERARCHY


def get_kipc_nodes() -> dict[int, dict[str, dict[str, Any]]]:
    """Return the official KIPC classification by year, keyed by KIPC code."""
    return _KIPC_NODES


# --------------------------------------------------------------------------
# HTTP
# --------------------------------------------------------------------------


def build_ssl_context() -> ssl.SSLContext:
    """Return a default-verifying context that also trusts the Russian state CA.

    ``rosstat.gov.ru`` serves a certificate issued by the Russian national CA,
    which no default trust store carries, and it does not send the issuing
    intermediate. Certificate verification and hostname checking stay enabled;
    the bundle only adds the two public certificates that chain the official
    host, and requests never leave ``ALLOWED_HOSTS``.
    """
    context = ssl.create_default_context()
    if CA_BUNDLE_PATH.exists():
        context.load_verify_locations(cafile=str(CA_BUNDLE_PATH))
    else:  # pragma: no cover - a deployment that removed the bundle
        logger.warning("Rosstat CA bundle missing at %s; primary host may fail", CA_BUNDLE_PATH)
    return context


def build_client() -> httpx.Client:
    """Build the single managed HTTP client used by one collection call."""
    return httpx.Client(
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": USER_AGENT, "Accept-Encoding": "gzip, deflate"},
        trust_env=True,
        follow_redirects=False,
        verify=build_ssl_context(),
    )


RETRYABLE_STATUSES: Final = frozenset({429, 500, 502, 503, 504})
RATE_LIMITED_STATUS: Final = 429


def _assert_allowed(url: str) -> None:
    """Refuse any URL outside the official Rosstat hosts."""
    host = urlsplit(url).hostname or ""
    if host not in ALLOWED_HOSTS:
        raise ValueError(f"Refusing to request a host outside the Rosstat allowlist: {host}")


def _retry_delay(attempt: int, response: httpx.Response | None) -> float:
    """Return the bounded wait before the next attempt."""
    delay = float(BACKOFF_FACTOR**attempt)
    if response is not None and response.status_code == RATE_LIMITED_STATUS:
        header = response.headers.get("Retry-After", "")
        delay = float(header) if header.isdigit() else RATE_LIMIT_BACKOFF
    return min(delay, MAX_RETRY_DELAY)


def fetch(client: httpx.Client, url: str) -> bytes:
    """GET one official URL with bounded retries, returning the decoded body."""
    _assert_allowed(url)
    last_error: Exception | None = None
    for attempt in range(MAX_RETRIES):
        try:
            response = client.get(url)
        except httpx.HTTPError as error:
            last_error = error
            logger.warning("Transport error for %s (attempt %d): %s", url, attempt + 1, error)
        else:
            if response.status_code == 200:
                content = response.content
                if len(content) > MAX_DOWNLOAD_BYTES:
                    raise ValueError(f"{url} returned {len(content)} bytes, above the ceiling")
                return content
            if response.status_code not in RETRYABLE_STATUSES:
                raise httpx.HTTPStatusError(
                    f"{url} returned HTTP {response.status_code}",
                    request=response.request,
                    response=response,
                )
            last_error = httpx.HTTPStatusError(
                f"{url} returned HTTP {response.status_code}",
                request=response.request,
                response=response,
            )
            logger.warning("Retryable HTTP %d for %s", response.status_code, url)
            time.sleep(_retry_delay(attempt, response))
            continue
        time.sleep(_retry_delay(attempt, None))
    raise RuntimeError(f"Could not retrieve {url} after {MAX_RETRIES} attempts") from last_error


def fetch_from_any_host(client: httpx.Client, path: str) -> tuple[bytes, str]:
    """Fetch one mediabank path, falling back to the mirrored English host."""
    errors: list[str] = []
    for host in HOSTS:
        url = urljoin(host, path)
        try:
            return fetch(client, url), url
        except (RuntimeError, httpx.HTTPError, ValueError) as error:
            errors.append(f"{host}: {error}")
            logger.warning("Host %s could not serve %s", host, path)
    raise RuntimeError(f"No official host served {path} ({'; '.join(errors)})")


# --------------------------------------------------------------------------
# Discovery
# --------------------------------------------------------------------------

_HREF = re.compile(r'href="([^"]+)"', re.IGNORECASE)


def discover_files(
    client: httpx.Client, section_url: str, patterns: dict[str, re.Pattern[str]]
) -> dict[str, str]:
    """Return the mediabank path of every wanted file linked from a section page.

    Monthly workbook names embed their reference month, so the current name is
    read from the official page rather than reconstructed.
    """
    _assert_allowed(section_url)
    html = fetch(client, section_url).decode("utf-8", errors="replace")
    candidates = [href for href in _HREF.findall(html) if MEDIABANK_PREFIX in href]
    found: dict[str, str] = {}
    for key, pattern in patterns.items():
        matches = sorted({href for href in candidates if pattern.search(href)})
        if not matches:
            continue
        # Later reference months and later end years sort last under the
        # zero-padded MM-YYYY and YYYY-YYYY spellings once the year leads.
        found[key] = max(matches, key=_recency_key)
    logger.info("Discovered %d of %d source files on %s", len(found), len(patterns), section_url)
    return found


def _recency_key(href: str) -> tuple[int, ...]:
    """Order discovered names newest last using the dates inside the filename."""
    name = href.rsplit("/", 1)[-1]
    month_year = re.search(r"_(\d{2})-(\d{4})\.", name)
    if month_year:
        return (int(month_year.group(2)), int(month_year.group(1)))
    years = [int(value) for value in re.findall(r"(\d{4})", name)]
    return (max(years) if years else 0, 0)


# --------------------------------------------------------------------------
# Workbook helpers
# --------------------------------------------------------------------------


def _text(value: Any) -> str:
    """Collapse a spreadsheet cell into comparable single-spaced text."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def _number(value: Any) -> float | None:
    """Return a finite float for a published cell, or None for a gap marker."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = _text(value)
    if not text:
        return None
    # Rosstat footnote markers travel inside the cell, e.g. "104,672)".
    text = re.sub(r"\d\)$", "", text).replace(" ", "").replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _month_from_text(text: str) -> date | None:
    """Return the first day of the Russian-language month named in ``text``."""
    match = re.search(r"([А-Яа-яЁё]+)\s+(\d{4})", text)
    if not match:
        return None
    month = RUSSIAN_MONTHS.get(match.group(1).lower())
    return date(int(match.group(2)), month, 1) if month else None


def _publish_date(workbook: Any, label: str) -> None:
    """Record the 'Обновлено' stamp the contents sheet carries, when present."""
    if "Содержание" not in workbook.sheetnames:
        return
    sheet = workbook["Содержание"]
    cells = [_text(cell) for row in sheet.iter_rows(max_col=6, values_only=True) for cell in row]
    for index, value in enumerate(cells):
        if "бновлено" in value:
            # The stamp sits in the cell directly below the label, which is one
            # full row further along this flattened scan.
            for candidate in cells[index : index + 14]:
                stamped = _month_day_year(candidate)
                if stamped:
                    _LAST_PUBLISH_DATE[label] = stamped
                    return
    logger.warning("No publication stamp found in %s", label)


def _month_day_year(text: str) -> date | None:
    """Parse Rosstat's '11 сентября 2026 г.' publication stamp."""
    match = re.search(r"(\d{1,2})\s+([А-Яа-яЁё]+)\s+(\d{4})", text)
    if not match:
        return None
    month = RUSSIAN_MONTHS.get(match.group(2).lower())
    if not month:
        return None
    return date(int(match.group(3)), month, int(match.group(1)))


def _open_workbook(payload: bytes, *, styled: bool = False) -> Any:
    """Open a downloaded workbook, optionally keeping the cell styles."""
    return load_workbook(io.BytesIO(payload), data_only=True, read_only=not styled, rich_text=False)


def _header_row(
    sheet: Worksheet, needles: tuple[str, ...], limit: int = 14
) -> tuple[int, dict[str, int]]:
    """Locate the header row and the column index of each wanted header.

    Each needle must match a *distinct* cell of the row. Matching the row's
    concatenated text instead would select the workbook title, which spells out
    the same words ("СТРУКТУРА ПОТРЕБИТЕЛЬСКИХ РАСХОДОВ НАСЕЛЕНИЯ (СПР)...")
    inside a single merged cell and carries no columns.
    """
    for index, row in enumerate(sheet.iter_rows(max_row=limit, max_col=20, values_only=True), 1):
        texts = [_text(cell).lower() for cell in row]
        # A header cell labels a column; the title is a sentence spanning the sheet.
        headers = {text: position for position, text in enumerate(texts) if 0 < len(text) <= 120}
        if all(any(needle in text for text in headers) for needle in needles):
            return index, headers
    raise ValueError(f"Header row not found in sheet {sheet.title!r}")


def _column_for(columns: dict[str, int], needle: str) -> int | None:
    """Return the position of the first header containing ``needle``."""
    for text, position in columns.items():
        if needle in text:
            return position
    return None


# --------------------------------------------------------------------------
# ipc_spr - monthly index and weights for every published node
# --------------------------------------------------------------------------

_DETAIL_MEASURE_COLUMNS: Final = (
    ("предыдущему месяцу", "MOM"),
    ("декабрю предыдущего года", "YTD"),
    ("соответствующему месяцу предыдущего года", "YOY"),
)


def _detail_sheet_rows(sheet: Worksheet, styled_sheet: Worksheet | None) -> list[dict[str, Any]]:
    """Return one parsed row per published node of a monthly detail sheet."""
    rows: list[dict[str, Any]] = []
    for index, row in enumerate(sheet.iter_rows(min_row=1, max_col=6, values_only=True), 1):
        name, code = _text(row[0]), _text(row[1])
        if not name or not code:
            continue
        if not re.match(r"^[0-9]", code):
            continue
        bold = False
        if styled_sheet is not None:
            font = styled_sheet.cell(index, 1).font
            bold = bool(font and font.bold)
        rows.append(
            {
                "row": index,
                "name": name,
                "code": code,
                "weight": _number(row[2]),
                "MOM": _number(row[3]),
                "YTD": _number(row[4]),
                "YOY": _number(row[5]),
                "is_group": bold,
            }
        )
    return rows


def _detail_level(code: str, is_group: bool) -> str:
    """Return the published hierarchy level of one Rosstat CPI node."""
    if code == HEADLINE_CODE:
        return "HEADLINE"
    if code in AGGREGATE_CODES:
        return "AGGREGATE"
    return "GROUP" if is_group else "ITEM"


def parse_monthly_detail(
    payload: bytes, source_url: str
) -> tuple[dict[date, dict[str, float]], dict[date, dict[str, float]]]:
    """Parse ipc_spr into ``(observations, weights)`` keyed by reference month.

    The weight column is the officially published consumer-expenditure structure
    in percent of the national total; it is stored exactly as published.
    """
    values_workbook = _open_workbook(payload)
    styled_workbook = _open_workbook(payload, styled=True)
    _publish_date(values_workbook, "monthly_detail")
    observations: dict[date, dict[str, float]] = {}
    weights: dict[date, dict[str, float]] = {}
    for sheet_name in values_workbook.sheetnames:
        if sheet_name == "Содержание":
            continue
        sheet = values_workbook[sheet_name]
        header_text = " ".join(
            _text(cell)
            for row in sheet.iter_rows(max_row=5, max_col=3, values_only=True)
            for cell in row
        )
        reference = _month_from_text(header_text)
        if reference is None:
            logger.warning("Skipping detail sheet %r: no reference month in its title", sheet_name)
            continue
        rows = _detail_sheet_rows(sheet, styled_workbook[sheet_name])
        if not rows:
            logger.warning("Skipping detail sheet %r: no data rows", sheet_name)
            continue
        month_values: dict[str, float] = {}
        month_weights: dict[str, float] = {}
        for entry in rows:
            level = _detail_level(entry["code"], entry["is_group"])
            node = normalise_node(entry["code"])
            for measure in ("MOM", "YTD", "YOY"):
                value = entry[measure]
                if value is None:
                    continue
                series_id = build_series_id("RSTG", level, entry["code"], measure)
                month_values[series_id] = value
                _register_series(series_id, entry, level, measure, source_url)
            if entry["weight"] is not None:
                primary = build_series_id("RSTG", level, entry["code"], "MOM")
                month_weights[primary] = entry["weight"]
            record = _HIERARCHY.setdefault(
                node,
                {
                    "node": node,
                    "native_code": entry["code"],
                    "name": entry["name"],
                    "level": level,
                    "parent": _tree_parent(entry["code"]) or "",
                    "kipc_code": "",
                    "kipc_parent": "",
                },
            )
            # Sheets are read oldest to newest, so the newest basket wins.
            record["weight"] = entry["weight"]
            record["name"] = entry["name"]
        observations[reference] = month_values
        weights[reference] = month_weights
    values_workbook.close()
    styled_workbook.close()
    logger.info("Monthly detail workbook: %d reference months parsed", len(observations))
    return observations, weights


def _tree_parent(code: str) -> str | None:
    """Return the parent local code for the published aggregation tree."""
    for parent, children in TREE_PARENTS.items():
        if code in children:
            return parent
    return None


def _register_series(
    series_id: str, entry: dict[str, Any], level: str, measure: str, source_url: str
) -> None:
    """Record the upstream descriptive row metadata.py will standardize."""
    UPSTREAM_METADATA.setdefault(
        series_id,
        {
            "native_code": entry["code"],
            "native_name": entry["name"],
            "level": level,
            "measure": measure,
            "classification": "RSTG",
            "source_url": source_url,
            "release_name": RELEASE_NAMES["monthly_detail"],
            "parent_code": _tree_parent(entry["code"]) or "",
        },
    )


# --------------------------------------------------------------------------
# ipc_mes - month-on-month index for the four headline aggregates since 1991
# --------------------------------------------------------------------------

_MOM_BLOCK_MARKER: Final = "к концу предыдущего месяца"


def parse_monthly_headline(payload: bytes, source_url: str) -> dict[date, dict[str, float]]:
    """Parse ipc_mes into month-on-month observations for the four aggregates."""
    workbook = _open_workbook(payload)
    _publish_date(workbook, "monthly_headline")
    observations: dict[date, dict[str, float]] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name == "Содержание":
            continue
        sheet = workbook[sheet_name]
        grid = [list(row) for row in sheet.iter_rows(max_col=48, values_only=True)]
        title = _text(grid[0][0]) if grid else ""
        code = next((value for needle, value in HEADLINE_SHEET_CONCEPTS if needle in title), None)
        if code is None:
            logger.warning("Skipping headline sheet %r: unrecognised title %r", sheet_name, title)
            continue
        years, year_row = _headline_years(grid)
        if not years:
            logger.warning("Skipping headline sheet %r: no year header", sheet_name)
            continue
        name, level = _headline_descriptor(code)
        series_id = build_series_id("RSTG", level, code, "MOM")
        UPSTREAM_METADATA.setdefault(
            series_id,
            {
                "native_code": code,
                "native_name": name,
                "level": level,
                "measure": "MOM",
                "classification": "RSTG",
                "source_url": source_url,
                "release_name": RELEASE_NAMES["monthly_headline"],
                "parent_code": _tree_parent(code) or "",
            },
        )
        count = _read_headline_block(grid, year_row, years, series_id, observations)
        logger.info("Headline sheet %r (%s): %d monthly observations", sheet_name, code, count)
    workbook.close()
    return observations


def _headline_descriptor(code: str) -> tuple[str, str]:
    """Return the official Russian name and level of one headline aggregate."""
    names = {
        "1": "Все товары и услуги",
        "6": "Продовольственные товары",
        "7": "Непродовольственные товары",
        "9000": "Услуги",
    }
    return names[code], ("HEADLINE" if code == HEADLINE_CODE else "AGGREGATE")


def _headline_years(grid: list[list[Any]]) -> tuple[dict[int, int], int]:
    """Locate the year header row and map each year to its column index."""
    for index, row in enumerate(grid[:12]):
        years = {
            int(_text(cell)): position
            for position, cell in enumerate(row)
            if re.fullmatch(r"(19|20)\d{2}", _text(cell))
        }
        if len(years) >= 5:
            return years, index
    return {}, -1


def _read_headline_block(
    grid: list[list[Any]],
    year_row: int,
    years: dict[int, int],
    series_id: str,
    observations: dict[date, dict[str, float]],
) -> int:
    """Read the month-on-month block that follows the year header."""
    count = 0
    in_block = False
    for row in grid[year_row + 1 :]:
        label = _text(row[0]).lower()
        if _MOM_BLOCK_MARKER in label:
            in_block = True
            continue
        month = RUSSIAN_MONTHS.get(label)
        if month is None:
            # The workbook continues with the December-on-December block; the
            # month labels repeat there, so the first non-month label ends it.
            if in_block and label:
                break
            continue
        if not in_block:
            continue
        for year, column in years.items():
            value = _number(row[column]) if column < len(row) else None
            if value is None:
                continue
            observations.setdefault(date(year, month, 1), {})[series_id] = value
            count += 1
    return count


# --------------------------------------------------------------------------
# SDMX - official headline index levels
# --------------------------------------------------------------------------

_SDMX_SERIES = re.compile(r"<Series\b([^>]*)>(.*?)</Series>", re.DOTALL)
_SDMX_OBS = re.compile(r"<Obs\b([^>]*)/?>")
_SDMX_ATTR = re.compile(r'(\w+)="([^"]*)"')


def parse_sdmx(payload: bytes, measure: str, source_url: str) -> dict[date, dict[str, float]]:
    """Parse one SDMX 2.1 StructureSpecificData message into CPI index levels.

    Only the CPI data domain is read: the SDDS message that carries CPI also
    carries the producer price index, which this collector does not model.
    """
    document = payload.decode("utf-8", errors="replace")
    observations: dict[date, dict[str, float]] = {}
    for header, body in _SDMX_SERIES.findall(document):
        attributes = dict(_SDMX_ATTR.findall(header))
        if attributes.get("DATA_DOMAIN") != "CPI" or attributes.get("REF_AREA") != "RU":
            continue
        if attributes.get("FREQ") != "M":
            logger.warning("Skipping SDMX series with frequency %r", attributes.get("FREQ"))
            continue
        series_id = build_series_id("SDDS", "HEADLINE", "ALL", measure)
        UPSTREAM_METADATA.setdefault(
            series_id,
            {
                "native_code": attributes.get("INDICATOR", "PCPI_IX"),
                "native_name": "Consumer price index, all items",
                "level": "HEADLINE",
                "measure": measure,
                "classification": "SDDS",
                "source_url": source_url,
                "release_name": RELEASE_NAMES["sdmx"],
                "parent_code": "",
                "base_period": attributes.get("BASE_PER", SDMX_BASE_PERIODS[measure]),
            },
        )
        for observation in _SDMX_OBS.findall(body):
            fields = dict(_SDMX_ATTR.findall(observation))
            period, raw = fields.get("TIME_PERIOD", ""), fields.get("OBS_VALUE", "")
            match = re.fullmatch(r"(\d{4})-(\d{2})", period)
            value = _number(raw)
            if not match or value is None:
                logger.warning("Skipping SDMX observation %r=%r", period, raw)
                continue
            reference = date(int(match.group(1)), int(match.group(2)), 1)
            observations.setdefault(reference, {})[series_id] = value
    return observations


# --------------------------------------------------------------------------
# Annual consumer-expenditure structure - basket history before 2025
# --------------------------------------------------------------------------

_SHEET_YEAR = re.compile(r"(20\d{2})")
_EMBEDDED_CODE = re.compile(r"^\s*([0-9]+(?:\.[А-Яа-яA-Za-z]+)?)\s+(\S.*)$")


def _sheet_month_range(sheet_name: str) -> tuple[int, int, int] | None:
    """Return ``(year, first_month, last_month)`` a CPR sheet governs."""
    match = _SHEET_YEAR.search(sheet_name)
    if not match:
        return None
    year = int(match.group(1))
    # Rosstat split 2022 into two published baskets at 1 April.
    if "до 01.04" in sheet_name:
        return year, 1, 3
    if "с 01.04" in sheet_name:
        return year, 4, 12
    return year, 1, 12


def parse_annual_weights(payload: bytes, label: str) -> dict[date, dict[str, float]]:
    """Parse a CPR workbook into monthly weight rows keyed by local code.

    Rosstat publishes one consumer-expenditure structure per calendar year (two
    in 2022) and uses it for every month of that year, so the annual basket is
    repeated over the months it actually governs rather than transformed.
    """
    workbook = _open_workbook(payload)
    _publish_date(workbook, label)
    weights: dict[date, dict[str, float]] = {}
    skipped: list[str] = []
    for sheet_name in workbook.sheetnames:
        if sheet_name == "Содержание":
            continue
        span = _sheet_month_range(sheet_name)
        if span is None:
            skipped.append(sheet_name)
            continue
        try:
            rows = _annual_weight_rows(workbook[sheet_name])
        except ValueError as error:
            logger.warning("Skipping %s sheet %r: %s", label, sheet_name, error)
            continue
        if not rows:
            skipped.append(sheet_name)
            continue
        year, first_month, last_month = span
        for month in range(first_month, last_month + 1):
            weights.setdefault(date(year, month, 1), {}).update(rows)
    if skipped:
        logger.info("%s: %d sheets carry no local code and were skipped", label, len(skipped))
    workbook.close()
    return weights


def _annual_weight_rows(sheet: Worksheet) -> dict[str, float]:
    """Return ``{local_code: weight}`` for one annual basket sheet."""
    header_index, columns = _header_row(sheet, ("спр",))
    code_column = _column_for(columns, "код")
    if code_column is None:
        raise ValueError("no published local code column")
    weight_column = _column_for(columns, "спр")
    name_column = _column_for(columns, "наименование")
    embedded = name_column is None or name_column == code_column
    rows: dict[str, float] = {}
    for row in sheet.iter_rows(min_row=header_index + 1, max_col=12, values_only=True):
        if weight_column is None or weight_column >= len(row):
            continue
        weight = _number(row[weight_column])
        raw = _text(row[code_column]) if code_column < len(row) else ""
        if weight is None or not raw:
            continue
        code = raw
        if embedded:
            match = _EMBEDDED_CODE.match(raw)
            if not match:
                continue
            code = match.group(1)
        if not re.fullmatch(r"[0-9]+(\.[А-Яа-яA-Za-z]+)?", code):
            continue
        rows[code] = weight
    return rows


# --------------------------------------------------------------------------
# KIPC - the official COICOP classification, used for hierarchy and audit
# --------------------------------------------------------------------------


# Rosstat pads a skipped classification level with a zero segment, so
# ``07.1.1.0.0`` is the level-3 node ``07.1.1``.
def canonical_kipc(code: str) -> str:
    """Return a KIPC code with its zero placeholder segments removed."""
    parts = code.split(".")
    while len(parts) > 1 and parts[-1] == "0":
        parts.pop()
    return ".".join(parts)


def kipc_parent(code: str, known: set[str]) -> str | None:
    """Return the nearest published ancestor of one KIPC code."""
    parts = canonical_kipc(code).split(".")
    while len(parts) > 1:
        parts.pop()
        candidate = canonical_kipc(".".join(parts))
        if candidate in known:
            return candidate
    return None


# Analytical KIPC aggregates that repeat the whole basket under a separate code
# and must not be summed with the twelve/thirteen divisions.
KIPC_ANALYTICAL_PREFIX: Final = "16"
_KIPC_CODE = re.compile(r"^\d{2}(\.\d+)*$")


def parse_kipc_weights(payload: bytes, label: str) -> dict[int, dict[str, dict[str, Any]]]:
    """Parse the KIPC workbook into ``{year: {code: {name, weight}}}``."""
    workbook = _open_workbook(payload)
    _publish_date(workbook, label)
    years: dict[int, dict[str, dict[str, Any]]] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name == "Содержание":
            continue
        sheet = workbook[sheet_name]
        try:
            header_index, columns = _header_row(sheet, ("кипц",))
        except ValueError as error:
            logger.warning("Skipping KIPC sheet %r: %s", sheet_name, error)
            continue
        header = next(
            iter(
                sheet.iter_rows(
                    min_row=header_index, max_row=header_index, max_col=20, values_only=True
                )
            )
        )
        year_columns = {
            int(_text(cell)): position
            for position, cell in enumerate(header)
            if re.fullmatch(r"20\d{2}", _text(cell))
        }
        if not year_columns:
            span = _sheet_month_range(sheet_name)
            if span is None:
                logger.warning("Skipping KIPC sheet %r: no year column", sheet_name)
                continue
            weight_column = _column_for(columns, "спр")
            if weight_column is None:
                logger.warning("Skipping KIPC sheet %r: no weight column", sheet_name)
                continue
            year_columns = {span[0]: weight_column}
        code_column = _column_for(columns, "код") or 0
        name_column = _column_for(columns, "наименование")
        sheet_years: dict[int, dict[str, dict[str, Any]]] = {}
        duplicates = 0
        for row in sheet.iter_rows(min_row=header_index + 1, max_col=20, values_only=True):
            code = canonical_kipc(_text(row[code_column]) if code_column < len(row) else "")
            if not _KIPC_CODE.match(code) or code.split(".")[0] == KIPC_ANALYTICAL_PREFIX:
                continue
            name = (
                _text(row[name_column])
                if name_column is not None and name_column < len(row)
                else ""
            )
            for year, column in year_columns.items():
                weight = _number(row[column]) if column < len(row) else None
                if weight is None:
                    continue
                nodes = sheet_years.setdefault(year, {})
                if code in nodes:
                    # Rosstat occasionally prints a class and its only member
                    # under one KIPC code. Adding the two would double the node,
                    # so the first (higher) row defines it.
                    duplicates += 1
                    continue
                nodes[code] = {"name": name, "weight": weight}
        if duplicates:
            logger.info(
                "KIPC sheet %r repeats %d codes; the first row wins", sheet_name, duplicates
            )
        # 2022 appears both in the multi-year sheet and in its own post-1 April
        # sheet. The later sheet supersedes the earlier regime for that year.
        years.update(sheet_years)
    workbook.close()
    logger.info("KIPC classification: %d published years", len(years))
    return years


# --------------------------------------------------------------------------
# Public collection contract
# --------------------------------------------------------------------------


def _reset_caches() -> None:
    """Clear every per-run cache so a rerun never reuses stale descriptions."""
    UPSTREAM_METADATA.clear()
    _SOURCE_FILES.clear()
    _LAST_PUBLISH_DATE.clear()
    _HIERARCHY.clear()
    _KIPC_NODES.clear()


class RosstatPublication(NamedTuple):
    """Everything one collection call read from the official publications."""

    observations: dict[date, dict[str, float | None]]
    weights: dict[date, dict[str, float]]
    kipc: dict[int, dict[str, dict[str, Any]]]
    files: dict[str, str]
    # The two surfaces that both publish the four headline month-on-month
    # aggregates are kept apart so validation can compare them on the overlap.
    detail_panel: dict[date, dict[str, float]]
    headline_panel: dict[date, dict[str, float]]


def collect_publication(start_date: date | None = None) -> RosstatPublication:
    """Download and parse every official Rosstat CPI surface this collector uses."""
    _reset_caches()
    observations: dict[date, dict[str, float | None]] = {}
    weights: dict[date, dict[str, float]] = {}
    kipc: dict[int, dict[str, dict[str, Any]]] = {}
    with build_client() as client:
        files = discover_files(client, PRICE_SECTION_URL, _FILE_PATTERNS)
        files.update(discover_files(client, SDMX_SECTION_URL, _SDMX_PATTERNS))
        wanted = set(_FILE_PATTERNS) | set(_SDMX_PATTERNS)
        missing = sorted(wanted - set(files))
        required = {"monthly_detail", "monthly_headline"}
        absent = required - set(files)
        if absent:
            raise RuntimeError(f"Official section page no longer links {sorted(absent)}")
        if missing:
            logger.warning("Optional source files not linked this release: %s", missing)

        detail_values, detail_weights = _read(client, files, "monthly_detail", parse_monthly_detail)
        observations = _merge(observations, detail_values)
        weights = _merge(weights, detail_weights)

        headline = _read(client, files, "monthly_headline", parse_monthly_headline)
        # The combined publication is the richer surface, so it wins wherever the
        # two overlap; ipc_mes only extends the four aggregates back to 1991.
        observations = _merge(observations, headline, overwrite=False)

        for key, measure in SDMX_MEASURES.items():
            if key not in files:
                continue
            payload, url = _download(client, files[key])
            observations = _merge(observations, parse_sdmx(payload, measure, url))

        for key in ("item_weights", "group_weights"):
            if key not in files:
                continue
            payload, url = _download(client, files[key])
            _SOURCE_FILES[key] = url
            annual = _weights_to_series(parse_annual_weights(payload, key), key)
            weights = _merge(weights, annual, overwrite=False)

        if "kipc_weights" in files:
            payload, url = _download(client, files["kipc_weights"])
            _SOURCE_FILES["kipc_weights"] = url
            kipc = parse_kipc_weights(payload, "kipc_weights")
            _KIPC_NODES.update(kipc)
            _annotate_kipc_codes(kipc)

    # The annual basket sheets cover a whole calendar year, so the current year
    # yields weight rows for months Rosstat has not published an index for yet.
    # Weights are only kept up to the newest published reference month.
    if observations:
        horizon = max(observations)
        weights = {key: value for key, value in weights.items() if key <= horizon}
    if start_date is not None:
        observations = {key: value for key, value in observations.items() if key >= start_date}
        weights = {key: value for key, value in weights.items() if key >= start_date}
        detail_values = {key: value for key, value in detail_values.items() if key >= start_date}
        headline = {key: value for key, value in headline.items() if key >= start_date}
        # A rewound run must describe exactly the series it carries: a basket
        # position retired before the window has no observation in it, and a
        # weight for such a series would reference a series this run never saw.
        kept = {series_id for values in observations.values() for series_id in values}
        for series_id in set(UPSTREAM_METADATA) - kept:
            del UPSTREAM_METADATA[series_id]
        weights = {
            reference: {
                series_id: weight for series_id, weight in values.items() if series_id in kept
            }
            for reference, values in weights.items()
        }
    logger.info(
        "Collected %d observation months and %d weight months across %d series",
        len(observations),
        len(weights),
        len(UPSTREAM_METADATA),
    )
    return RosstatPublication(
        observations, weights, kipc, dict(_SOURCE_FILES), detail_values, headline
    )


def _weights_to_series(
    by_code: dict[date, dict[str, float]], label: str
) -> dict[date, dict[str, float]]:
    """Key annual basket weights by the node's primary series identifier.

    The annual publications carry only the local code, so the level comes from
    the monthly publication that was parsed first. A code with no monthly series
    is a basket position Rosstat has since retired; it is dropped rather than
    stored against a series that does not exist.
    """
    resolved: dict[date, dict[str, float]] = {}
    unknown: set[str] = set()
    for reference, values in by_code.items():
        bucket = resolved.setdefault(reference, {})
        for code, weight in values.items():
            node = _HIERARCHY.get(normalise_node(code))
            if node is None:
                unknown.add(code)
                continue
            bucket[build_series_id("RSTG", node["level"], code, "MOM")] = weight
    if unknown:
        logger.info(
            "%s: %d basket codes have no current monthly series and were skipped",
            label,
            len(unknown),
        )
    return resolved


def _annotate_kipc_codes(kipc: dict[int, dict[str, dict[str, Any]]]) -> int:
    """Attach the official KIPC code to every item Rosstat publishes in both.

    Rosstat prints the representative goods and services under the same official
    names in the monthly workbook and in the KIPC classification, and publishes
    the identical basket weight for them. Matching on that published name is
    therefore a lookup of Rosstat's own cross-listing, not a guess; the weight is
    compared as a second key so a renamed or re-weighted item stays unannotated.
    """
    if not kipc:
        return 0
    year = max(kipc)
    by_name: dict[str, list[tuple[str, float]]] = {}
    for code, node in kipc[year].items():
        by_name.setdefault(node["name"].casefold(), []).append((code, float(node["weight"])))
    matched = 0
    for node in _HIERARCHY.values():
        candidates = by_name.get(node["name"].casefold(), [])
        if len(candidates) != 1:
            continue
        code, weight = candidates[0]
        if node.get("weight") is not None and abs(float(node["weight"]) - weight) > 1e-9:
            continue
        node["kipc_code"] = code
        node["kipc_parent"] = kipc_parent(code, set(kipc[year])) or ""
        matched += 1
    for fields in UPSTREAM_METADATA.values():
        described = _HIERARCHY.get(normalise_node(str(fields.get("native_code", ""))))
        if described and described.get("kipc_code"):
            fields["kipc_code"] = described["kipc_code"]
    logger.info(
        "KIPC cross-listing: %d of %d nodes carry an official KIPC code", matched, len(_HIERARCHY)
    )
    return matched


def _download(client: httpx.Client, path: str) -> tuple[bytes, str]:
    """Fetch one discovered mediabank path, pausing between files."""
    if DOWNLOAD_DELAY:
        time.sleep(DOWNLOAD_DELAY)
    return fetch_from_any_host(client, path)


def _read(client: httpx.Client, files: dict[str, str], key: str, parser: Any) -> Any:
    """Download and parse one required workbook, recording its resolved URL."""
    payload, url = _download(client, files[key])
    _SOURCE_FILES[key] = url
    return parser(payload, url)


def _merge(
    target: dict[date, dict[str, Any]], extra: dict[date, dict[str, Any]], *, overwrite: bool = True
) -> dict[date, dict[str, Any]]:
    """Merge one parsed panel into another, keeping the preferred source."""
    for reference, values in extra.items():
        bucket = target.setdefault(reference, {})
        for series_id, value in values.items():
            if overwrite or series_id not in bucket:
                bucket[series_id] = value
    return target


def collect_raw_data(start_date: date | None = None) -> dict[date, dict[str, float | None]]:
    """Return ``{reference_date: {series_id: value}}`` for every collected series."""
    return collect_publication(start_date).observations
