"""Write the analyst-facing audit workbook for one collected publication.

The workbook mirrors what the database holds so an analyst can reproduce every
published aggregate by hand: the official observations, the official basket used
to reconstruct them, the published hierarchy including Rosstat's KIPC
cross-listing, and the reconciliation result of every check the run performed.
It is a validation aid and never a substitute for the standardized tables.
"""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from scripts.validate import Check

logger = logging.getLogger(__name__)

_HIERARCHY_HEADERS = (
    "series_node",
    "native_code",
    "name",
    "level",
    "rosstat_parent_code",
    "kipc_code",
    "kipc_parent_code",
    "latest_weight_percent",
)
_CHECK_HEADERS = (
    "kind",
    "reference",
    "node",
    "measure",
    "official",
    "reconstructed",
    "difference",
    "tolerance",
    "result",
)


def _panel_sheet(workbook: Workbook, title: str, panel: dict[date, dict[str, Any]]) -> None:
    """Write one date-by-series panel, dates as rows and series as columns."""
    sheet = workbook.create_sheet(title)
    columns = sorted({series_id for values in panel.values() for series_id in values})
    sheet.append(["reference_date", *columns])
    for reference in sorted(panel):
        values = panel[reference]
        sheet.append([reference.isoformat(), *(values.get(column) for column in columns)])
    logger.info("Audit workbook %s: %d rows x %d series", title, len(panel), len(columns))


def export_validation_xlsx(
    path: Path,
    observations: dict[date, dict[str, float | None]],
    weights: dict[date, dict[str, float]],
    hierarchy: dict[str, dict[str, Any]],
    kipc: dict[int, dict[str, dict[str, Any]]],
    checks: list[Check],
) -> Path:
    """Write the audit workbook and return the path it was written to."""
    workbook = Workbook()
    # A new workbook always opens with one empty default sheet; the audit sheets
    # are created explicitly below.
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    _panel_sheet(workbook, "Time Series", observations)
    _panel_sheet(workbook, "Weights", weights)

    sheet = workbook.create_sheet("Hierarchy")
    sheet.append(list(_HIERARCHY_HEADERS))
    for node in sorted(hierarchy.values(), key=lambda row: str(row["native_code"])):
        sheet.append(
            [
                node["node"],
                node["native_code"],
                node["name"],
                node["level"],
                node.get("parent", ""),
                node.get("kipc_code", ""),
                node.get("kipc_parent", ""),
                node.get("weight"),
            ]
        )

    sheet = workbook.create_sheet("KIPC Basket")
    sheet.append(["year", "kipc_code", "name", "weight_percent"])
    for year in sorted(kipc):
        for code in sorted(kipc[year]):
            node = kipc[year][code]
            sheet.append([year, code, node["name"], node["weight"]])

    sheet = workbook.create_sheet("Validation")
    sheet.append(list(_CHECK_HEADERS))
    for check in checks:
        sheet.append(list(check))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    logger.info("Audit workbook written to %s", path)
    return path
